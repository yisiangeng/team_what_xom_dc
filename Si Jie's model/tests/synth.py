"""Synthetic raw files 01 and 02 with the real layout (spec task1 section 4).

Truth:
- 5 quarterly history dates (GRID). Rates are piecewise constant on each
  quarter (t_q, t_{q+1}], so cumulatives are piecewise linear with C(t_0) = 0.
- Cases 1 to 4 report on the grid dates plus extra-date set A, case 5 on
  set B and case 6 on set C (two extra dates in each quarter, one with a time
  of day). The Date column is the union, with empty cells.
- Case 2 has exactly the rates of RAMP_History (V4 test A09).
- RAMP_History reports only on the grid dates. It produces from t_1 on. Its
  reported rate is the next-quarter mean (or the previous-quarter mean when
  history_convention='prev'); at t_{T-1} the value is arbitrary.
- 'Field Production' copies the reported history rates.
- 'Oil production rate' ends with one all-empty row.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import yaml

GRID = [
    dt.datetime(2001, 1, 1),
    dt.datetime(2001, 4, 1),
    dt.datetime(2001, 7, 1),
    dt.datetime(2001, 10, 1),
    dt.datetime(2002, 1, 1),
]
T = len(GRID)
CASES = [1, 2, 3, 4, 5, 6]
HIST_CASE = 2  # case whose rates equal the history's
FLUIDS = ("oil", "gas", "water")

EXTRA_SETS = {
    "A": [dt.timedelta(days=20), dt.timedelta(days=50, hours=8, minutes=41, milliseconds=938)],
    "B": [dt.timedelta(days=30), dt.timedelta(days=61, hours=13, minutes=5, seconds=7, milliseconds=123)],
    "C": [dt.timedelta(days=10), dt.timedelta(days=70, hours=21, minutes=17, seconds=45, milliseconds=555)],
}
CASE_SET = {1: "A", 2: "A", 3: "A", 4: "A", 5: "B", 6: "C"}

# quantity key -> (sheet name, unit, fluid, kind)
QUANT = {
    "oil_rate": ("Oil production rate", "STB/d", "oil", "rate"),
    "oil_cum": ("Oil production cumulative", "STB", "oil", "cum"),
    "gas_rate": ("Gas production rate", "MSCF/d", "gas", "rate"),
    "gas_cum": ("Gas production cumulative", "MSCF", "gas", "cum"),
    "water_rate": ("Water production rate", "STB/d", "water", "rate"),
    "water_cum": ("Water production cumulative", "STB", "water", "cum"),
}
# Sheet order of the real file 02
SHEET_ORDER = [
    "Gas production cumulative",
    "Gas production rate",
    "Oil production cumulative",
    "Oil production rate",
    "Water production cumulative",
    "Water production rate",
]

HIST_RATES = {
    "oil": [0.0, 1500.0, 1400.0, 1200.0],
    "gas": [0.0, 1200.0, 1260.0, 1200.0],
    "water": [0.0, 0.0, 100.0, 300.0],
}
HIST_LAST_REPORTED = {"oil": 1111.0, "gas": 999.0, "water": 333.0}  # arbitrary at t_{T-1}

PARAMS = {  # case -> (fault_trans, poro_mult, perm_mult, aquifer_pv)
    1: (0.05, 0.80, 0.60, 50.0),
    2: (0.12, 1.20, 3.00, 200.0),
    3: (0.08, 1.00, 9.50, 120.0),
    4: (0.15, 0.90, 1.50, 80.0),
    5: (0.02, 1.40, 5.00, 150.0),
    6: (0.10, 1.10, 2.20, 60.0),
}
FILE01_ROW_ORDER = [1, 2, 3, 4, 6, 5]  # not sorted on purpose
FILE01_HEADER = ["Item", "Fault Transmissibility", "Porosity Multiplier",
                 "Permeability Multiplier", "Aquifer Pore Volume"]


def case_rates(i: int) -> dict[str, list[float]]:
    """Piecewise-constant rates of case i on quarters q = 0..T-2."""
    if i == HIST_CASE:
        return {f: list(v) for f, v in HIST_RATES.items()}
    oil = [800.0 + 100.0 * i, 1600.0 - 20.0 * i, 1300.0 + 10.0 * i, 1100.0 - 5.0 * i]
    gas = [oil[q] * (0.7 + 0.05 * i + 0.1 * q) for q in range(T - 1)]
    water = [0.0, 50.0 * i, 150.0 + 20.0 * i, 300.0 + 30.0 * i]
    return {"oil": oil, "gas": gas, "water": water}


def days(d: dt.datetime) -> float:
    """Days from GRID[0], computed like the code: integer ms / 86,400,000."""
    ms = (d - GRID[0]) // dt.timedelta(milliseconds=1)
    return ms / 86_400_000


GRID_DAYS = [days(g) for g in GRID]


def cum_at(rates: list[float], tau: float) -> float:
    total = 0.0
    for q in range(T - 1):
        lo, hi = GRID_DAYS[q], GRID_DAYS[q + 1]
        overlap = min(tau, hi) - lo
        if overlap > 0:
            total += rates[q] * overlap
    return total


def rate_at(rates: list[float], tau: float) -> float:
    """Reported (simulator last-timestep) rate at native time tau."""
    if tau <= GRID_DAYS[0]:
        return 0.0
    for q in range(T - 1):
        if GRID_DAYS[q] < tau <= GRID_DAYS[q + 1]:
            return rates[q]
    raise ValueError(tau)


def case_dates(i: int) -> list[dt.datetime]:
    extra = [GRID[q] + off for q in range(T - 1) for off in EXTRA_SETS[CASE_SET[i]]]
    return sorted(GRID + extra)


def reported_history(convention: str) -> dict[str, list[float]]:
    out = {}
    for f in FLUIDS:
        r = HIST_RATES[f]
        if convention == "next":
            out[f] = [r[k] for k in range(T - 1)] + [HIST_LAST_REPORTED[f]]
        elif convention == "prev":
            out[f] = [0.0] + [r[k - 1] for k in range(1, T)]
        else:
            raise ValueError(convention)
    return out


def synthetic_config(root: Path) -> dict:
    root = Path(root)
    return {
        "seed": 0,
        "paths": {
            "file01": str(root / "raw" / "01 Introduction.xlsx"),
            "file02": str(root / "raw" / "02 Field Production.xlsx"),
            "db": str(root / "processed" / "dwc26.sqlite"),
            "arrays": str(root / "processed" / "arrays.npz"),
            "report": str(root / "outputs" / "reports" / "data_checks.md"),
            "logs": str(root / "outputs" / "logs"),
        },
        "file01": {
            "sheet": "Uncertainty Values",
            "label_column": "Item",
            "label_pattern": r"^Case\s+(\d+)$",
            "columns": {
                "Fault Transmissibility": "fault_trans",
                "Porosity Multiplier": "poro_mult",
                "Permeability Multiplier": "perm_mult",
                "Aquifer Pore Volume": "aquifer_pv",
            },
        },
        "file02": {
            "names_row": 2,
            "units_row": 3,
            "first_data_row": 4,
            "date_column": "Date",
            "case_pattern": r"^Case_(\d+)$",
            "history_column": "RAMP_History",
            "sheets": {q: {"name": v[0], "unit": v[1]} for q, v in QUANT.items()},
            "field_sheet": {
                "name": "Field Production",
                "columns": {
                    "Date": "date",
                    "Oil production rate (STB/d)": "oil_rate",
                    "Gas production rate (MSCF/d)": "gas_rate",
                    "Water production rate (STB/d)": "water_rate",
                    "Water injection rate (STB/d)": "water_inj_rate",
                },
            },
            "date_round": "ms",
        },
        "split": {"train": [1, 4], "val": [5, 5], "test": [6, 6]},
        "x_bounds": {"source": "design", "override": None, "citation": None},
        "rates": {"definition": "auto"},
        "checks": {
            "c4_rel_tol": 1.0e-6,
            "c6_rel_tol": 0.01,
            "c6_floor_frac": 1.0e-3,
            "c8_abs_tol": 0.0,
            "c10_rel_tol": 1.0e-5,
        },
    }


def write_config(cfg: dict, path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(cfg, sort_keys=False), encoding="utf-8")
    return path


@dataclass
class SynthRaw:
    file01: Path
    file02: Path
    cfg: dict
    root: Path
    n_nonempty: int = 0
    reported: dict = field(default_factory=dict)

    def config_path(self, name: str = "data.yaml") -> Path:
        return write_config(self.cfg, self.root / name)


def _empty_styled_row(ws, ncols: int) -> None:
    r = ws.max_row + 1
    for c in range(1, ncols + 1):
        ws.cell(row=r, column=c).number_format = "0.00"


def write_file01(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "Nomenclature"
    ws0.append(["Parameter", "Meaning"])
    ws0.append(["Fault Transmissibility", "multiplier"])
    ws = wb.create_sheet("Uncertainty Values")
    ws.append(FILE01_HEADER)
    for i in FILE01_ROW_ORDER:
        ws.append([f"Case {i}", *PARAMS[i]])
    _empty_styled_row(ws, len(FILE01_HEADER))
    wb.save(path)


def write_file02(path: Path, convention: str) -> tuple[int, dict]:
    all_dates = sorted(set(d for i in CASES for d in case_dates(i)) | set(GRID))
    reported = reported_history(convention)
    wb = openpyxl.Workbook()
    wb.active.title = "Plots"
    fp = wb.create_sheet("Field Production")
    fp.append(["Identifier", "Date", "Oil production rate (STB/d)", "Gas production rate (MSCF/d)",
               "Water production rate (STB/d)", "Water injection rate (STB/d)"])
    for k, g in enumerate(GRID):
        fp.append(["Field", g, reported["oil"][k], reported["gas"][k], reported["water"][k], 2000.0 * k])

    n_nonempty = 0
    by_name = {v[0]: (q, v) for q, v in QUANT.items()}
    names = ["Date"] + [f"Case_{i}" for i in CASES] + ["RAMP_History"]
    for sheet in SHEET_ORDER:
        q, (sname, unit, fluid, kind) = by_name[sheet]
        ws = wb.create_sheet(sheet)
        ws.append(["X"] + [f"Y{j}" for j in range(1, len(names))])
        ws.append(names)
        ws.append([None] + [f"{sname} [{unit}]"] * (len(names) - 1))
        dates_of = {i: set(case_dates(i)) for i in CASES}
        for d in all_dates:
            tau = days(d)
            row = [d]
            for i in CASES:
                if d in dates_of[i]:
                    r = case_rates(i)[fluid]
                    row.append(cum_at(r, tau) if kind == "cum" else rate_at(r, tau))
                    n_nonempty += 1
                else:
                    row.append(None)
            if d in GRID:
                k = GRID.index(d)
                if kind == "cum":
                    row.append(cum_at(HIST_RATES[fluid], tau))
                else:
                    row.append(reported[fluid][k])
                n_nonempty += 1
            else:
                row.append(None)
            ws.append(row)
        if sheet == "Oil production rate":
            _empty_styled_row(ws, len(names))
    wb.save(path)
    return n_nonempty, reported


def build_raw(root: Path, *, history_convention: str = "next") -> SynthRaw:
    root = Path(root)
    cfg = synthetic_config(root)
    p1, p2 = Path(cfg["paths"]["file01"]), Path(cfg["paths"]["file02"])
    p1.parent.mkdir(parents=True, exist_ok=True)
    write_file01(p1)
    n, reported = write_file02(p2, history_convention)
    return SynthRaw(file01=p1, file02=p2, cfg=cfg, root=root, n_nonempty=n, reported=reported)


# --- helpers to corrupt copies -------------------------------------------------

def edit_xlsx(path: Path, fn) -> None:
    wb = openpyxl.load_workbook(path)
    fn(wb)
    wb.save(path)


def col_of(ws, name: str, row: int = 2) -> int:
    for cell in ws[row]:
        if cell.value == name:
            return cell.column
    raise KeyError(name)


def row_of(ws, when: dt.datetime, first_data_row: int = 4, col: int = 1) -> int:
    for r in range(first_data_row, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == when:
            return r
    raise KeyError(when)
