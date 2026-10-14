"""Excel reading of files 01 and 02 (spec task1 sections 1.1 and 1.3).

The layout is read from the config and verified by check C1; nothing about
the layout is assumed beyond the configured rows and names. An empty cell
means "no report on this date"; it is never imputed (PRD section 5).
"""

from __future__ import annotations

import datetime as dt
import math
import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from dwc26.data.check import ContractError, assert_array

HISTORY_ID = 0          # case_id given to RAMP_History in long tables
QUANTITIES = ("oil_rate", "oil_cum", "gas_rate", "gas_cum", "water_rate", "water_cum")
FLUIDS = ("oil", "gas", "water")

_UNIT_RX = re.compile(r"\[([^\]]*)\]")


@dataclass
class RawProduction:
    sheets_found: list[str]
    names: dict[str, list[str]]         # quantity -> row-2 names after the Date column
    units: dict[str, list[str]]         # quantity -> unit text inside [...] for each data column ('' if none)
    date_col: dict[str, np.ndarray]     # quantity -> Date column, datetime64[ms], all-empty rows dropped
    native: pd.DataFrame                # case_id int64 (0 = RAMP_History), quantity, date datetime64[ms], value float64
    bad_cells: pd.DataFrame             # non-empty cells that do not parse to a finite float
    field: pd.DataFrame | None          # date, oil_rate, gas_rate, water_rate, water_inj_rate


# --- cell helpers --------------------------------------------------------------

def _is_empty(v) -> bool:
    return v is None or (isinstance(v, str) and v == "")


def _to_float(v) -> float | None:
    """Finite float, or None if the cell does not parse to a finite float."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        x = float(v)
    elif isinstance(v, str):
        try:
            x = float(v.strip())
        except ValueError:
            return None
    else:
        return None
    return x if math.isfinite(x) else None


def _name(v) -> str:
    return "" if v is None else str(v)


def _ts(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return pd.Timestamp(v)
    if isinstance(v, str):
        try:
            return pd.Timestamp(v)
        except (ValueError, TypeError):
            return pd.NaT
    return pd.NaT


def parse_dates(values, date_round: str) -> np.ndarray:
    """Cells -> datetime64[<date_round>], rounded to that unit (spec 2.2). Unparseable -> NaT."""
    idx = pd.DatetimeIndex([_ts(v) for v in values])
    if idx.tz is not None:
        idx = idx.tz_localize(None)
    return idx.round(date_round).to_numpy().astype(f"datetime64[{date_round}]")


def _open(path: Path):
    return openpyxl.load_workbook(Path(path), read_only=True, data_only=True)


def _sheet_rows(wb, name: str) -> list[tuple]:
    return list(wb[name].iter_rows(values_only=True))


def _pad(rows: list[tuple]) -> np.ndarray:
    width = max((len(r) for r in rows), default=0)
    arr = np.empty((len(rows), width), dtype=object)
    arr[:] = None
    for i, r in enumerate(rows):
        arr[i, : len(r)] = r
    return arr


# --- file 01 -------------------------------------------------------------------

def load_parameters(path: Path, cfg01: dict) -> tuple[pd.DataFrame, list[str]]:
    """Columns: case_id int64, then the internal names in config order (float64).
    Rows are sorted by case_id, and all-empty rows are dropped.
    The list holds structural problems (missing sheet or column, bad label,
    duplicate id, non-finite value); rows with a problem are not in the frame.
    C1 reports them. Raises only if the file cannot be opened."""
    names = list(cfg01["columns"].values())
    problems: list[str] = []

    def frame(records):
        df = pd.DataFrame.from_records(records, columns=["case_id", *names])
        df["case_id"] = df["case_id"].astype(np.int64)
        for n in names:
            df[n] = df[n].astype(np.float64)
        df = df.sort_values("case_id", kind="mergesort").reset_index(drop=True)
        # V9 boundary: rows with a problem are excluded, so this always holds
        ids = df["case_id"].to_numpy()
        assert_array("load_parameters.case_id", ids, (len(df),), dtype=np.int64)
        if len(np.unique(ids)) != len(ids):
            raise ContractError("load_parameters: case_id not unique")
        assert_array("load_parameters.values", df[names].to_numpy(dtype=np.float64), (len(df), len(names)))
        return df

    wb = _open(path)
    try:
        sheet = cfg01["sheet"]
        if sheet not in wb.sheetnames:
            problems.append(f"file 01: sheet '{sheet}' not found (sheets: {wb.sheetnames})")
            return frame([]), problems
        rows = _sheet_rows(wb, sheet)
    finally:
        wb.close()

    if not rows:
        problems.append(f"file 01: sheet '{sheet}' is empty")
        return frame([]), problems
    header = [_name(v) for v in rows[0]]
    col_idx: dict[str, int] = {}
    for col in [cfg01["label_column"], *cfg01["columns"]]:
        hits = [j for j, h in enumerate(header) if h == col]
        if not hits:
            problems.append(f"file 01: column '{col}' not found in the header {header}")
        elif len(hits) > 1:
            problems.append(f"file 01: column '{col}' appears {len(hits)} times")
        else:
            col_idx[col] = hits[0]
    if problems:
        return frame([]), problems

    rx = re.compile(cfg01["label_pattern"])
    seen: set[int] = set()
    records = []
    for r, row in enumerate(rows[1:], start=2):
        if all(_is_empty(v) for v in row):
            continue
        label = row[col_idx[cfg01["label_column"]]] if col_idx[cfg01["label_column"]] < len(row) else None
        if _is_empty(label):
            problems.append(f"file 01 row {r}: empty label")
            continue
        m = rx.fullmatch(_name(label))
        if not m:
            problems.append(f"file 01 row {r}: label '{label}' does not match {cfg01['label_pattern']}")
            continue
        cid = int(m.group(1))
        if cid in seen:
            problems.append(f"file 01 row {r}: duplicate case id {cid}")
            continue
        vals = []
        for col in cfg01["columns"]:
            j = col_idx[col]
            x = _to_float(row[j] if j < len(row) else None)
            if x is None:
                problems.append(f"file 01 row {r} ({label}): '{col}' = {row[j] if j < len(row) else None!r} "
                                "is not a finite number")
            vals.append(x)
        if any(x is None for x in vals):
            continue
        seen.add(cid)
        records.append((cid, *vals))
    return frame(records), problems


# --- file 02 -------------------------------------------------------------------

def _parse_data_sheet(rows, q, cfg02, case_rx, out):
    nr, ur, fr = cfg02["names_row"] - 1, cfg02["units_row"] - 1, cfg02["first_data_row"] - 1
    if len(rows) <= max(nr, ur):
        out["names"][q] = []
        return
    arr = _pad(rows)
    names_row, units_row, data = arr[nr], arr[ur], arr[fr:]
    # drop phantom trailing columns (no name, no unit, no data)
    width = arr.shape[1]
    while (width > 1 and _is_empty(names_row[width - 1]) and _is_empty(units_row[width - 1])
           and all(_is_empty(v) for v in data[:, width - 1])):
        width -= 1
    names_row, units_row, data = names_row[:width], units_row[:width], data[:, :width]
    # drop all-empty data rows
    if len(data):
        empty = pd.isna(data) | (data == "")
        data = data[~empty.all(axis=1)]

    if width == 0 or _name(names_row[0]) != cfg02["date_column"]:
        out["names"][q] = [_name(v) for v in names_row]  # Date column missing: C1 fails
        return
    names = [_name(v) for v in names_row[1:]]
    out["names"][q] = names
    units = []
    for u in units_row[1:]:
        s = _name(u)
        m = _UNIT_RX.search(s)
        units.append(m.group(1) if m else s.strip())
    out["units"][q] = units
    dates = parse_dates(data[:, 0], cfg02["date_round"]) if len(data) else np.array([], dtype=f"datetime64[{cfg02['date_round']}]")
    out["date_col"][q] = dates

    block = data[:, 1:]
    nonempty = ~pd.isna(block)
    for j, name in enumerate(names):
        m = case_rx.fullmatch(name)
        if m:
            cid = int(m.group(1))
        elif name == cfg02["history_column"]:
            cid = HISTORY_ID
        else:
            continue
        rows_j = np.flatnonzero(nonempty[:, j])
        vals, ok_rows, bad_rows, bad_raw = [], [], [], []
        for i in rows_j:
            v = block[i, j]
            if _is_empty(v):
                continue
            x = _to_float(v)
            if x is None:
                bad_rows.append(i)
                bad_raw.append(str(v))
            else:
                ok_rows.append(i)
                vals.append(x)
        if ok_rows:
            out["native"].append(pd.DataFrame({
                "case_id": np.full(len(ok_rows), cid, dtype=np.int64),
                "quantity": q,
                "date": dates[ok_rows],
                "value": np.array(vals, dtype=np.float64),
            }))
        if bad_rows:
            out["bad"].append(pd.DataFrame({
                "case_id": np.full(len(bad_rows), cid, dtype=np.int64),
                "quantity": q,
                "date": dates[bad_rows],
                "raw": bad_raw,
            }))


def _parse_field_sheet(rows, cfg02) -> pd.DataFrame:
    cols = cfg02["field_sheet"]["columns"]
    unit = cfg02["date_round"]
    if not rows:
        return pd.DataFrame()
    header = [_name(v) for v in rows[0]]
    data = [r for r in rows[1:] if not all(_is_empty(v) for v in r)]
    out = {}
    for head, internal in cols.items():
        if head not in header:
            continue
        j = header.index(head)
        vals = [r[j] if j < len(r) else None for r in data]
        if internal == "date":
            out[internal] = parse_dates(vals, unit)
        else:
            out[internal] = np.array([np.nan if (x := _to_float(v)) is None else x for v in vals],
                                     dtype=np.float64)
    return pd.DataFrame(out)


def load_production(path: Path, cfg02: dict) -> RawProduction:
    """Never raises on a structural deviation (missing sheet or column, extra
    column, wrong unit); C1 detects those. Columns that match neither
    case_pattern nor history_column appear in `names` but not in `native`."""
    unit = cfg02["date_round"]
    case_rx = re.compile(cfg02["case_pattern"])
    out = {"names": {}, "units": {}, "date_col": {}, "native": [], "bad": []}
    field = None
    wb = _open(path)
    try:
        sheets_found = list(wb.sheetnames)
        for q in QUANTITIES:
            name = cfg02["sheets"][q]["name"]
            if name not in sheets_found:
                continue
            _parse_data_sheet(_sheet_rows(wb, name), q, cfg02, case_rx, out)
        fname = cfg02["field_sheet"]["name"]
        if fname in sheets_found:
            field = _parse_field_sheet(_sheet_rows(wb, fname), cfg02)
    finally:
        wb.close()

    if out["native"]:
        native = pd.concat(out["native"], ignore_index=True)
    else:
        native = pd.DataFrame({"case_id": pd.Series(dtype=np.int64), "quantity": pd.Series(dtype=object),
                               "date": pd.Series(dtype=f"datetime64[{unit}]"),
                               "value": pd.Series(dtype=np.float64)})
    native = native.sort_values(["case_id", "quantity", "date"], kind="mergesort").reset_index(drop=True)
    if out["bad"]:
        bad = pd.concat(out["bad"], ignore_index=True)
    else:
        bad = pd.DataFrame({"case_id": pd.Series(dtype=np.int64), "quantity": pd.Series(dtype=object),
                            "date": pd.Series(dtype=f"datetime64[{unit}]"), "raw": pd.Series(dtype=object)})
    return RawProduction(
        sheets_found=sheets_found,
        names=out["names"],
        units=out["units"],
        date_col=out["date_col"],
        native=native,
        bad_cells=bad,
        field=field,
    )


def native_series(raw: RawProduction) -> dict[tuple[int, str], tuple[np.ndarray, np.ndarray]]:
    """(case_id, quantity) -> (dates, values) of the non-empty cells, sorted by date."""
    out = {}
    nat = raw.native
    if len(nat) == 0:
        return out
    for (cid, q), g in nat.groupby(["case_id", "quantity"], sort=False):
        out[(int(cid), str(q))] = (g["date"].to_numpy(), g["value"].to_numpy(dtype=np.float64))
    return out
