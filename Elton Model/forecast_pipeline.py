"""History-matched forecast pipeline for DWC26 (cumulative domain).

Stages (run from ``cod/`` with the venv interpreter):

    python forecast_pipeline.py --check-grid        # validate params -> grid features mapping
    python forecast_pipeline.py --train-cumulative  # cumulative twin proxy -> proxy_outputs/cumulative_model.pt
    python history_match.py                         # search parameters vs observed history -> hm_* artifacts
    python forecast_pipeline.py --forecast          # 20-year extension of the matches + selection of 10
    python forecast_pipeline.py --deliverable       # workspace root / 10_Final_deliverables.xlsx

Why cumulative: notebook 1 decided "use cumulative values for matching" (reported rate
columns follow different conventions in the history and in the cases). Matching and
forecasting therefore work entirely on cumulative production; the rate proxy
(``final_model.pt``) remains the validated step-1 artifact.

Forecast principle: the provided simulator data ends at 2008-01-01, so the 20-year
extension cannot be learned from data. Instead each matched case's cumulative curve is
extended by the parametric family (Arps cumulative / Gompertz / rational) that wins a
backtest on the 100 known cases (fit on quarters 1-25, scored on quarters 26-41).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import time
from dataclasses import asdict
from pathlib import Path

import numpy as np
import pandas as pd
import torch
import torch.nn.functional as F

import proxy_model as pm
from proxy_model import CODE, DATA, OUT, GRID_FEATURES, PARAMS, PHASES, SEED, Config

CUM_MODEL = OUT / "cumulative_model.pt"
WORKSPACE = CODE.parent
DELIVERABLE = WORKSPACE / "10_Final_deliverables.xlsx"
TEMPLATE = CODE / "dataset_ori" / "09 Template Deliverable.xlsx"

# Zone sizes measured in notebook 1 section 5.1 (active mask identical in all 100 cases).
N_ACTIVE, N_ZONE_PERM, N_ZONE_AQ = 42_512, 11_575, 2_286
FZ_PERM = N_ZONE_PERM / N_ACTIVE

FAMILY_ORDER = ["arps_cum", "gompertz", "rational"]


# --------------------------------------------------------------------------- data

def load_cumulative_data():
    """``pm.load_data()`` but with the cumulative targets of curves_quarterly_cum.npz.

    The npz stores its third axis in template order [gas, oil, water] (notebook 1,
    ``QTY``), while everything in this pipeline and in proxy_model.py uses
    PHASES = [oil, gas, water]. The axis is reordered once, here; the assert pins the
    stored order so a future re-generation of the npz cannot pass silently.
    """
    data = pm.load_data()
    z = np.load(DATA / "curves_quarterly_cum.npz")
    assert list(z["quantities"]) == ["gas_cum", "oil_cum", "water_cum"]  # notebook 1 order
    order = [list(z["quantities"]).index(f"{p}_cum") for p in PHASES]    # -> [1, 0, 2]
    data["Y"] = z["Y"][:, :, order].astype(np.float64)
    data["y_hist"] = z["y_hist"][:, order].astype(np.float64)
    return data


def forecast_dates():
    """The 80 quarterly template dates (2008-04-01 to 2028-01-01)."""
    return pd.read_csv(DATA / "forecast_dates.csv")


# ------------------------------------------------- params -> grid features mapping

def _case_one():
    cases = pd.read_csv(DATA / "cases.csv")
    feats = pd.read_csv(DATA / "grid_case_features.csv")
    p1 = cases.loc[cases.case_id == 1, PARAMS].iloc[0]
    f1 = feats.loc[feats.case_id == 1, GRID_FEATURES].iloc[0]
    return p1, f1


def grid_features_analytic(theta):
    """Analytic grid summaries for any parameter set.

    Every case grid is case 1 scaled by the parameter ratios (verified in notebook 1
    section 5.1 to <= 3e-5): porosity everywhere, permeability inside its zone, pore
    volume inside the aquifer zone. The eight summaries are therefore closed-form
    functions of the three ratios, anchored on the case-1 features.
    """
    theta = np.atleast_2d(np.asarray(theta, dtype=np.float64))
    p1, f1 = _case_one()
    r_poro = theta[:, 1] / float(p1["poro_mult"])
    r_perm = theta[:, 2] / float(p1["perm_mult"])
    r_aq = theta[:, 3] / float(p1["aquifer_pv"])
    fz = FZ_PERM
    w_tx = fz * float(f1["mean_permx_zone"]) / float(f1["mean_permx"])
    return np.column_stack([
        r_poro * (f1["reservoir_pv"] + r_aq * f1["aquifer_pv_grid"]),  # total_pv
        r_poro * r_aq * f1["aquifer_pv_grid"],                         # aquifer_pv_grid
        r_poro * f1["reservoir_pv"],                                   # reservoir_pv
        r_poro * f1["mean_poro"],                                      # mean_poro
        f1["mean_permx"] + (r_perm - 1.0) * fz * f1["mean_permx_zone"],# mean_permx
        r_perm * f1["mean_permx_zone"],                                # mean_permx_zone
        f1["mean_tranx"] * (1.0 + (r_perm - 1.0) * w_tx),              # mean_tranx
        f1["mean_trany"] * (1.0 + (r_perm - 1.0) * w_tx),              # mean_trany
    ])


class GridMapper:
    """Parameters -> 8 grid features, validated on all 100 cases.

    Primary method is the analytic scaling above. Acceptance: every feature must be
    within 0.2 sigma of the stored ``grid_case_features.csv`` values, where sigma is
    the train-case standard deviation of the feature (the number the checkpoints store
    in ``grid_scaler``). Otherwise a linear fallback on
    ``[1, r_poro, r_perm, r_poro * r_aq]`` is fitted on the 70 train cases and must
    pass the same check on the 30 held-out cases; if it does not, the pipeline stops.
    """

    TOL = 0.2

    def __init__(self, data=None):
        self.data = load_cumulative_data() if data is None else data
        cases = self.data["cases"]
        theta = cases[PARAMS].to_numpy(dtype=np.float64)
        stored = self.data["grid"][GRID_FEATURES].to_numpy(dtype=np.float64)
        train = cases["split"].to_numpy() == "train"
        sigma = np.where(stored[train].std(axis=0) > 0, stored[train].std(axis=0), 1.0)

        pred = grid_features_analytic(theta)
        err = np.abs(pred - stored) / sigma
        self.max_error = err.max(axis=0)
        self.method = "analytic" if (self.max_error < self.TOL).all() else "fallback"
        self.max_error_holdout = None
        if self.method == "fallback":
            self._fit_fallback(theta, stored, train, sigma)
        self.report = {
            "method": self.method,
            "tolerance_sigma": self.TOL,
            "max_error_sigma_per_feature": {f: float(e) for f, e in zip(GRID_FEATURES, self.max_error)},
            "max_error_holdout_sigma": (None if self.max_error_holdout is None
                                        else float(self.max_error_holdout)),
        }

    @staticmethod
    def _design(theta):
        p1, _ = _case_one()
        r_poro = theta[:, 1] / float(p1["poro_mult"])
        r_perm = theta[:, 2] / float(p1["perm_mult"])
        r_aq = theta[:, 3] / float(p1["aquifer_pv"])
        return np.column_stack([np.ones(len(theta)), r_poro, r_perm, r_poro * r_aq])

    def _fit_fallback(self, theta, stored, train, sigma):
        A = self._design(theta)
        self.coef_, *_ = np.linalg.lstsq(A[train], stored[train], rcond=None)
        holdout = ~train
        err = np.abs(A[holdout] @ self.coef_ - stored[holdout]) / sigma
        self.max_error_holdout = err.max()
        if self.max_error_holdout >= self.TOL:
            raise RuntimeError(
                "grid feature mapping failed: analytic and fallback both exceed "
                f"{self.TOL} sigma (holdout max {self.max_error_holdout:.3g})")

    def transform(self, theta):
        theta = np.atleast_2d(np.asarray(theta, dtype=np.float64))
        if self.method == "analytic":
            return grid_features_analytic(theta)
        return self._design(theta) @ self.coef_


_MAPPER = None


def grid_mapper():
    global _MAPPER
    if _MAPPER is None:
        _MAPPER = GridMapper()
    return _MAPPER


# --------------------------------------------------------------- cumulative model

def _scaler_from_ckpt(ckpt):
    scaler = pm.PhaseScaler(ckpt["scaler"]["use_log"])
    scaler.mean_ = np.asarray(ckpt["scaler"]["mean"])
    scaler.std_ = np.asarray(ckpt["scaler"]["std"])
    scaler.base_curve_ = np.asarray(ckpt["scaler"]["base_curve"])
    return scaler


def load_cumulative_model():
    # weights_only=False: the checkpoint holds numpy arrays and the config dict, not
    # only tensors. Only load checkpoints trained in this workspace.
    ckpt = torch.load(CUM_MODEL, map_location="cpu", weights_only=False)
    cfg_data = dict(ckpt["config"])
    cfg_data["hidden"] = tuple(cfg_data["hidden"])
    ckpt["_cfg"] = Config(**cfg_data)
    return ckpt


def predict_cumulative(ckpt, theta, grid=None, batch=8192):
    """Cumulative curves (n, 41, 3) for raw parameter sets."""
    theta = np.atleast_2d(np.asarray(theta, dtype=np.float64))
    if grid is None:
        grid = grid_mapper().transform(theta)
    p_min = np.asarray(ckpt["parameter_min"], dtype=np.float64)
    p_range = np.asarray(ckpt["parameter_max"], dtype=np.float64) - p_min
    if np.any(p_range <= 0):
        raise ValueError("checkpoint has a zero-width parameter range")
    g_mean = np.asarray(ckpt["grid_scaler"]["mean"], dtype=np.float64)
    g_std = np.asarray(ckpt["grid_scaler"]["std"], dtype=np.float64)
    X = np.hstack([(theta - p_min) / p_range, (grid - g_mean) / g_std])
    scaler = _scaler_from_ckpt(ckpt)
    parts = []
    for s in range(0, len(X), batch):
        parts.append(pm.predict_ensemble(ckpt["states"], scaler, X[s:s + batch], ckpt["_cfg"]))
    return np.concatenate(parts, axis=0)


def train_cumulative(n_seeds=5, epochs=None):
    """Train the cumulative twin proxy (same architecture/protocol as step 1)."""
    torch.manual_seed(SEED)
    np.random.seed(SEED)
    data = load_cumulative_data()
    cfg = pm.load_best_config()
    print(f"cumulative proxy: {cfg.feature_set} hidden={cfg.hidden} wd={cfg.weight_decay:g}")

    states, scaler, epochs = pm.train_final(data, cfg, n_seeds=n_seeds, epochs=epochs)
    X = pm.build_inputs(data, cfg.feature_set)
    pred = pm.predict_ensemble(states, scaler, X, cfg)

    # baselines on the same cumulative targets
    pred_ridge = pm.baseline_ridge(data, cfg)
    pred_mean = pm.baseline_mean_curve(data)

    splits = data["cases"]["split"].to_numpy()
    rows = []
    for sp in ["train", "validation", "test"]:
        m = splits == sp
        rows += pm.metrics_table(data["Y"], pred, m, sp, "dnn_ensemble")
        rows += pm.metrics_table(data["Y"], pred_ridge, m, sp, "ridge")
        rows += pm.metrics_table(data["Y"], pred_mean, m, sp, "mean_curve")
    metrics = pd.DataFrame(rows)
    metrics.to_csv(OUT / "cumulative_metrics.csv", index=False)
    print(metrics.to_string(index=False, float_format=lambda v: f"{v:,.4g}"))

    # gate: beat the ridge baseline on the held-out test cases
    t = metrics[(metrics.split == "test") & (metrics.phase == "all")]
    dnn_rmse = float(t.loc[t.model == "dnn_ensemble", "rmse"].iloc[0])
    ridge_rmse = float(t.loc[t.model == "ridge", "rmse"].iloc[0])
    rel = {ph: float(metrics[(metrics.split == "test") & (metrics.phase == ph) &
                             (metrics.model == "dnn_ensemble")]["rmse"].iloc[0]) /
                float(data["y_hist"][-1][j]) for j, ph in enumerate(PHASES)}
    print("test RMSE as fraction of the observed final cumulative:",
          {k: f"{v:.2%}" for k, v in rel.items()})
    if dnn_rmse >= ridge_rmse:
        raise RuntimeError(f"gate failed: cumulative proxy test RMSE {dnn_rmse:,.3g} "
                           f"does not beat ridge {ridge_rmse:,.3g}")

    torch.save({
        "config": asdict(cfg) | {"hidden": list(cfg.hidden)},
        "seeds": [SEED + 100 + k for k in range(n_seeds)],
        "states": states,
        "scaler": {"use_log": scaler.use_log, "mean": scaler.mean_,
                   "std": scaler.std_, "base_curve": scaler.base_curve_},
        "phases": PHASES, "params": PARAMS, "grid_features": GRID_FEATURES,
        "grid_scaler": {"mean": data["grid"][GRID_FEATURES].to_numpy(dtype=np.float64)[
                            splits == "train"].mean(axis=0),
                        "std": np.where(
                            data["grid"][GRID_FEATURES].to_numpy(dtype=np.float64)[
                                splits == "train"].std(axis=0) > 0,
                            data["grid"][GRID_FEATURES].to_numpy(dtype=np.float64)[
                                splits == "train"].std(axis=0), 1.0)},
        "dates": np.datetime_as_string(data["dates"], unit="D"),
        "t_days": data["t_days"],
        "parameter_min": data["cases"].loc[splits == "train", PARAMS].min().to_numpy(dtype=float),
        "parameter_max": data["cases"].loc[splits == "train", PARAMS].max().to_numpy(dtype=float),
        "epochs": epochs,
        "env": {"torch": torch.__version__, "python": pm.platform.python_version()},
        "target": "cumulative",
    }, CUM_MODEL)
    print(f"wrote {CUM_MODEL}")
    return ckpt_summary()


def ckpt_summary():
    return {"path": str(CUM_MODEL), "sha256": hashlib.sha256(CUM_MODEL.read_bytes()).hexdigest()}


# ------------------------------------------------------------------ misfit (HM)

def hm_misfit_parts(pred, y_hist):
    """Normalized cumulative misfit; returns (total (n,), per phase (n, 3)).

    m = sqrt(mean over phases and dates of ((pred - obs) / obs_final)^2).
    """
    pred = np.asarray(pred, dtype=np.float64)
    y_hist = np.asarray(y_hist, dtype=np.float64)
    d = (pred - y_hist[None]) / y_hist[-1][None, None]
    per_phase = np.sqrt((d ** 2).mean(axis=1))
    return np.sqrt((d ** 2).mean(axis=(1, 2))), per_phase


def case_misfits(data):
    """Misfit of the 100 provided cases (same objective, on their stored curves)."""
    return hm_misfit_parts(data["Y"], data["y_hist"])


# ------------------------------------------------------------ curve families

class CurveFit:
    """Fitted parametric family for a batch of curves (raw params + anchors)."""

    def __init__(self, name, z, c_last, t_last):
        self.name, self.z, self.c_last, self.t_last = name, z, c_last, t_last


def _family_eval(name, z, t, c_last, t_last):
    """Evaluate family ``name`` for a batch. z (n, 3) raw, t (T, ) days, anchors per curve."""
    if name == "arps_cum":       # Arps hyperbolic rates -> closed-form cumulative
        q0 = (0.05 + F.softplus(z[:, 0])) * 4.0 * (c_last / t_last)
        Di = (0.05 + F.softplus(z[:, 1])) * 4.0 / t_last
        b = 0.2 + 2.0 * torch.sigmoid(z[:, 2])
        x = 1.0 + b[:, None] * Di[:, None] * t[None]
        a = (1.0 - 1.0 / b)[:, None]
        den = Di[:, None] * (b[:, None] - 1.0)
        den = torch.where(b[:, None] > 1.0, den.clamp(min=1e-12), den.clamp(max=-1e-12))
        arps = q0[:, None] * (x.pow(a) - 1.0) / den
        lin = q0[:, None] * torch.log1p(Di[:, None] * t[None]) / Di[:, None]
        near = (torch.abs(b - 1.0) < 1e-3)[:, None]
        return torch.where(near, lin, arps)
    if name == "gompertz":
        K = (0.2 + F.softplus(z[:, 0])) * 2.0 * c_last
        bb = (0.05 + F.softplus(z[:, 1])) * 4.0 / t_last
        t0 = z[:, 2] * 0.3 * t_last
        return K[:, None] * torch.exp(-torch.exp(-bb[:, None] * (t[None] - t0[:, None])))
    if name == "rational":
        A = (0.2 + F.softplus(z[:, 0])) * 2.0 * c_last
        p = 0.2 + 2.5 * torch.sigmoid(z[:, 1])
        B = (0.05 + F.softplus(z[:, 2])) * (0.5 * t_last) ** p
        tp = t[None].pow(p[:, None])
        return A[:, None] * tp / (tp + B[:, None])
    raise ValueError(name)


def fit_family(name, t_fit, C_fit, steps=1200, lr=0.05):
    """Fit one family to a batch of curves. C_fit: (n, T) cumulative values."""
    t = torch.as_tensor(np.asarray(t_fit, dtype=np.float32))
    C = torch.as_tensor(np.asarray(C_fit, dtype=np.float32))
    c_last = C[:, -1].clamp(min=1e-9)
    z = torch.zeros((C.shape[0], 3), requires_grad=True)
    opt = torch.optim.Adam([z], lr=lr)
    sched = torch.optim.lr_scheduler.CosineAnnealingLR(opt, T_max=steps, eta_min=lr * 0.1)
    t_last = float(t[-1])
    for _ in range(steps):
        opt.zero_grad()
        pred = _family_eval(name, z, t, c_last, t_last)
        loss = (((pred - C) / c_last[:, None]) ** 2).mean()
        loss.backward()
        opt.step()
        sched.step()
    with torch.no_grad():
        fit = CurveFit(name, z.detach(), c_last, t_last)
    return fit


def eval_fit(fit, t_new):
    """Evaluate a fitted family at new times -> (n, T_new) numpy array."""
    t = torch.as_tensor(np.asarray(t_new, dtype=np.float32))
    with torch.no_grad():
        return _family_eval(fit.name, fit.z, t, fit.c_last, fit.t_last).numpy()


def backtest_families(Y, t_days, n_fit=25, steps=1200):
    """Pick the family per phase by fitting on the first ``n_fit`` quarters of every
    provided case and scoring quarters ``n_fit``..41 against the stored truth."""
    t = np.asarray(t_days, dtype=float)
    rows = []
    for name in FAMILY_ORDER:
        for j, phase in enumerate(PHASES):
            fit = fit_family(name, t[:n_fit], Y[:, :n_fit, j], steps=steps)
            pred = eval_fit(fit, t[n_fit:])
            anchor = Y[:, n_fit - 1, j][:, None]        # cumulative at the last fitted quarter
            rel = np.sqrt((((pred - Y[:, n_fit:, j]) / anchor) ** 2).mean(axis=1))
            rows.append({"family": name, "phase": phase,
                         "median_relative_rmse": float(np.median(rel)),
                         "mean_relative_rmse": float(rel.mean()),
                         "p90_relative_rmse": float(np.percentile(rel, 90)),
                         "family_order": FAMILY_ORDER.index(name)})
    table = pd.DataFrame(rows)
    winners = {}
    for phase in PHASES:
        sub = table[table.phase == phase].sort_values(
            ["median_relative_rmse", "mean_relative_rmse", "family_order"])
        winners[phase] = str(sub.iloc[0]["family"])
    return table, winners


def forecast_from_candidates(C_hist, t_hist, t_future, winners, steps=1200):
    """Extend candidate cumulative curves (n, 41, 3) to the forecast dates.

    The winning family per phase is refit on all 41 quarters; the extension is shifted
    to stay continuous with the predicted history at 2008-01-01 and forced monotone.
    """
    C_hist = np.asarray(C_hist, dtype=np.float64)
    t_hist = np.asarray(t_hist, dtype=float)
    t_future = np.asarray(t_future, dtype=float)
    out = np.empty((C_hist.shape[0], len(t_future), len(PHASES)))
    for j, phase in enumerate(PHASES):
        fit = fit_family(winners[phase], t_hist, C_hist[:, :, j], steps=steps)
        fut = eval_fit(fit, t_future)
        anchor_pred = eval_fit(fit, t_hist[-1:])[:, 0]      # family value at 2008-01-01
        fut = fut + (C_hist[:, -1, j:j + 1] - anchor_pred[:, None])
        fut = np.maximum.accumulate(fut, axis=1)            # monotone guard
        out[:, :, j] = np.maximum(fut, C_hist[:, -1, j:j + 1])
    return out


# ------------------------------------------------------------------- selection

def select_top10(theta, score, k=10, seed=0):
    """One best-scoring representative per parameter-space cluster (diversity guard)."""
    from sklearn.cluster import KMeans

    theta = np.asarray(theta, dtype=np.float64)
    Z = (theta - theta.mean(axis=0)) / np.where(theta.std(axis=0) > 0, theta.std(axis=0), 1.0)
    labels = KMeans(n_clusters=min(k, len(theta)), random_state=seed, n_init=10).fit_predict(Z)
    best = [int(np.argmax(np.where(labels == c, score, -np.inf))) for c in np.unique(labels)]
    return np.array(sorted(best, key=lambda i: -score[i])), labels


def _spearman(a, b):
    ra = np.argsort(np.argsort(a))
    rb = np.argsort(np.argsort(b))
    return float(np.corrcoef(ra, rb)[0, 1])


# ----------------------------------------------------------------- deliverable

def write_deliverable(selected, out_path, template=TEMPLATE):
    """Fill the 10 template sheets and validate the result.

    ``selected``: list of dicts ``{"rank": 1..10, "gas": (80,), "oil": (80,), "water": (80,)}``
    ordered best-first. Returns a readback report; raises on any structural failure.
    """
    from openpyxl import load_workbook

    if len(selected) != 10:
        raise ValueError(f"need 10 cases, got {len(selected)}")
    wb = load_workbook(template)
    for item in selected:
        ws = wb[f"Case {int(item['rank'])}"]
        for k in range(80):
            for col, key in [(2, "gas"), (3, "oil"), (4, "water")]:
                cell = ws.cell(row=2 + k, column=col, value=float(np.round(item[key][k], 0)))
                cell.number_format = "0"
    wb.save(out_path)

    # readback validation against the template and the written values
    report = {"path": str(out_path), "sheets": {}}
    for item in selected:
        name = f"Case {int(item['rank'])}"
        tpl = pd.read_excel(template, sheet_name=name, header=None)
        out = pd.read_excel(out_path, sheet_name=name, header=None)
        dates_ok = (tpl.iloc[1:, 0].astype("string").fillna("").to_numpy()
                    == out.iloc[1:, 0].astype("string").fillna("").to_numpy()).all()
        vals = np.asarray(out.iloc[1:, 1:4], dtype=float)
        checks = {
            "rows": int(vals.shape[0]),
            "dates_match_template": bool(dates_ok),
            "finite": bool(np.isfinite(vals).all()),
            "monotone": bool((np.diff(vals, axis=0) >= -1e-6).all()),
        }
        ok = checks["rows"] == 80 and checks["dates_match_template"] and checks["finite"] \
            and checks["monotone"]
        report["sheets"][name] = checks | {"ok": ok}
        if not ok:
            raise RuntimeError(f"deliverable validation failed for {name}: {checks}")
    return report


# ------------------------------------------------------------------- stages

def run_check_grid():
    rep = grid_mapper().report
    print(json.dumps(rep, indent=2))
    return rep


def run_forecast(steps=1200):
    data = load_cumulative_data()
    ckpt = load_cumulative_model()
    t = np.asarray(data["t_days"], dtype=float)

    print("family backtest (fit quarters 1-25, score 26-41, all 100 cases) ...")
    table, winners = backtest_families(data["Y"], t, n_fit=25, steps=steps)
    table.to_csv(OUT / "family_backtest.csv", index=False)
    print(table.to_string(index=False, float_format=lambda v: f"{v:.4f}"))
    print("winning family per phase:", winners)

    cand = pd.read_csv(OUT / "hm_candidates.csv").sort_values("misfit").reset_index(drop=True)
    m_min = float(cand["misfit"].iloc[0])
    W = cand[cand["misfit"] <= 1.25 * m_min].head(100)
    if len(W) < 10:
        W = cand.head(20)
    print(f"well-matching set: {len(W)} candidates (m <= {1.25 * m_min:.4g})")

    theta_w = W[PARAMS].to_numpy(dtype=np.float64)
    C_w = predict_cumulative(ckpt, theta_w)
    fut = forecast_from_candidates(C_w, t, forecast_dates()["t_days"].to_numpy(float),
                                   winners, steps=steps)
    score = fut[:, -1, 0]                                   # oil cumulative at 2028
    sel, labels = select_top10(theta_w, score, k=10)
    print(f"spearman(misfit, oil_2028) in the well-matching set: {_spearman(W['misfit'].to_numpy(), score):.2f}")

    sel_rows, fc_rows = [], []
    fdates = forecast_dates()
    for rank, i in enumerate(sel, start=1):
        row = {"rank": rank, "cluster": int(labels[i]), "misfit": float(W["misfit"].iloc[i])}
        row |= {p: float(theta_w[i, j]) for j, p in enumerate(PARAMS)}
        for j, ph in enumerate(PHASES):
            row[f"{ph}_cum_2008"] = float(C_w[i, -1, j])
            row[f"{ph}_cum_2028"] = float(fut[i, -1, j])
        sel_rows.append(row)
        for k, (_, fd) in enumerate(fdates.iterrows()):
            for j, ph in enumerate(PHASES):
                fc_rows.append({"rank": rank, "date": fd["date"], "t_days": float(fd["t_days"]),
                                "phase": ph, "cum": float(fut[i, k, j])})
    sel_df = pd.DataFrame(sel_rows)
    sel_df.to_csv(OUT / "hm_top10_parameters.csv", index=False)
    pd.DataFrame(fc_rows).to_csv(OUT / "forecast_cum.csv", index=False)
    print(sel_df[["rank", "misfit", *PARAMS, "oil_cum_2028"]].to_string(index=False))

    # fan chart: observed history + selected candidates' predicted history and forecast
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, axes = plt.subplots(3, 1, figsize=(11, 9), sharex=True)
    hist_dates = pd.DatetimeIndex(data["dates"])
    fut_dates = pd.DatetimeIndex(fdates["date"])
    colors = plt.cm.tab10(np.linspace(0, 1, 10))
    for j, ph in enumerate(PHASES):
        a = axes[j]
        a.plot(hist_dates, data["y_hist"][:, j], color="black", lw=2.4, label="observed history")
        for rank, i in enumerate(sel, start=1):
            a.plot(hist_dates, C_w[i, :, j], color=colors[rank - 1], lw=1.1, alpha=0.9)
            a.plot(fut_dates, fut[i, :, j], color=colors[rank - 1], lw=1.4,
                   label=f"case {rank}" if j == 0 else None)
        a.set_title(f"{ph} cumulative")
        a.grid(alpha=0.3)
    axes[0].legend(fontsize=8, ncol=4)
    fig.suptitle("Selected 10 cases: predicted history (41 q) and forecast to 2028 (80 q)")
    fig.tight_layout()
    fig.savefig(OUT / "forecast_fan.png", dpi=150)
    plt.close(fig)
    print(f"wrote {OUT / 'forecast_cum.csv'}, {OUT / 'hm_top10_parameters.csv'}, {OUT / 'forecast_fan.png'}")
    return sel_df, winners


def run_deliverable():
    sel_df = pd.read_csv(OUT / "hm_top10_parameters.csv")
    fc = pd.read_csv(OUT / "forecast_cum.csv")
    selected = []
    for _, row in sel_df.iterrows():
        item = {"rank": int(row["rank"])}
        for ph in PHASES:
            sub = fc[(fc["rank"] == row["rank"]) & (fc["phase"] == ph)].sort_values("t_days")
            item[ph] = sub["cum"].to_numpy(dtype=float)
            if len(item[ph]) != 80:
                raise RuntimeError(f"forecast_cum.csv has {len(item[ph])} rows for rank {row['rank']} {ph}")
        # continuity with the predicted 2008-01-01 cumulative
        for ph in PHASES:
            if item[ph][0] < float(row[f"{ph}_cum_2008"]) - 1e-6:
                raise RuntimeError(f"rank {row['rank']} {ph} starts below its 2008 anchor")
        selected.append(item)
    report = write_deliverable(selected, DELIVERABLE)
    print(json.dumps(report, indent=2))
    print(f"\nwrote {DELIVERABLE}")
    return report


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check-grid", action="store_true", help="validate params -> grid features mapping")
    ap.add_argument("--train-cumulative", action="store_true", help="train the cumulative twin proxy")
    ap.add_argument("--forecast", action="store_true", help="extend history matches to 2028 and select 10")
    ap.add_argument("--deliverable", action="store_true", help="write 10_Final_deliverables.xlsx")
    ap.add_argument("--seeds", type=int, default=5)
    ap.add_argument("--epochs", type=int, default=None)
    args = ap.parse_args()

    t0 = time.time()
    did = False
    if args.check_grid:
        run_check_grid()
        did = True
    if args.train_cumulative:
        train_cumulative(n_seeds=args.seeds, epochs=args.epochs)
        did = True
    if args.forecast:
        run_forecast()
        did = True
    if args.deliverable:
        run_deliverable()
        did = True
    if not did:
        ap.print_help()
    else:
        print(f"done in {time.time() - t0:.0f} s")


if __name__ == "__main__":
    main()
